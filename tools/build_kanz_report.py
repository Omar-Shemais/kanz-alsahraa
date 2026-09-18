import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report-source.md"
OUTPUT = ROOT / "تقرير_الصيانة_الشامل_لتطبيق_كنز_الصحراء_2026.docx"

BLACK = "000000"
GOLD = "7A5C19"
PALE_GOLD = "F7F2E7"
PALE_GRAY = "F5F5F5"
LIGHT_BORDER = "D9D9D9"
FONT = "Arial"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), LIGHT_BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, name=FONT, size=None, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.28


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GOLD)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    r_pr.extend([color, underline, rtl])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, size=11.5, color=BLACK):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(size - 0.5, 8), color="333333")
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            add_hyperlink(paragraph, link.group(1), link.group(2))
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    specs = {
        "Title": (24, True, 18, 10),
        "Subtitle": (13, False, 8, 6),
        "Heading 1": (17, True, 14, 7),
        "Heading 2": (13.5, True, 10, 5),
        "Heading 3": (12, True, 8, 4),
    }
    for name, (size, bold, before, after) in specs.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style.element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)

    list_style = styles["List Bullet"]
    list_style.font.name = FONT
    list_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    list_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    list_style._element.rPr.rFonts.set(qn("w:cs"), FONT)
    list_style.font.size = Pt(11.2)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def setup_footer(doc):
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.paragraph_format.space_after = Pt(0)
        table = footer.add_table(rows=1, cols=2, width=Inches(6.8))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(1.0)
        table.columns[1].width = Inches(5.8)
        page_p = table.cell(0, 0).paragraphs[0]
        page_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(page_p)
        title_p = table.cell(0, 1).paragraphs[0]
        set_paragraph_rtl(title_p, WD_ALIGN_PARAGRAPH.RIGHT)
        title_p.paragraph_format.space_after = Pt(0)
        run = title_p.add_run("كنز الصحراء  |  تقرير الصيانة")
        set_run_font(run, size=8.5, color="666666")


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    set_paragraph_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
    title_pr = title._p.get_or_add_pPr()
    title_border = title_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_pr.remove(title_border)
    add_inline(title, "العرض الفني والتجاري لصيانة وتطوير تطبيق كنز الصحراء", size=24)
    subtitle = doc.add_paragraph()
    set_paragraph_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    add_inline(subtitle, "فحص الكود والتكامل وخطة الإصلاح والاختبار والنشر", size=13)
    doc.add_paragraph()
    details = [
        "مقدم إلى إدارة كنز الصحراء",
        "إعداد م عمر",
        "٨ سبتمبر ٢٠٢٦",
        "عرض مقترح للتنفيذ",
    ]
    for text in details:
        p = doc.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, text, size=11.5, color="333333")
    for _ in range(4):
        doc.add_paragraph()
    note = doc.add_paragraph()
    set_paragraph_rtl(note, WD_ALIGN_PARAGRAPH.CENTER)
    add_inline(note, "وثيقة سرية للاستخدام الإداري والفني", size=9, color="666666")
    doc.add_page_break()


def parse_table(lines, start):
    block = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        block.append(lines[idx].strip())
        idx += 1
    rows = []
    for line in block:
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows, idx


def add_table(doc, rows):
    if not rows:
        return
    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = {
        2: [2.0, 4.7],
        3: [1.55, 2.15, 3.0],
        4: [1.2, 2.0, 1.7, 1.8],
    }.get(columns, [6.7 / columns] * columns)
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(columns):
            cell = row.cells[c_idx]
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, GOLD if r_idx == 0 else (PALE_GOLD if r_idx % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, values[c_idx] if c_idx < len(values) else "", size=9.2, color="FFFFFF" if r_idx == 0 else BLACK)
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown_body(doc, markdown):
    lines = markdown.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## القرار التنفيذي"))
    i = start
    paragraph_lines = []

    def flush_paragraph():
        nonlocal paragraph_lines
        if paragraph_lines:
            p = doc.add_paragraph()
            set_paragraph_rtl(p)
            add_inline(p, " ".join(s.strip() for s in paragraph_lines), size=11.5)
            paragraph_lines = []

    page_break_before = {"مراجعة التقرير السابق", "خطة التنفيذ", "العرض المالي"}
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped.startswith("### ") or stripped.startswith("## "):
            flush_paragraph()
            level = 2 if stripped.startswith("### ") else 1
            text = stripped[4:] if level == 2 else stripped[3:]
            if text in page_break_before:
                doc.add_page_break()
            p = doc.add_paragraph(style=f"Heading {level}")
            set_paragraph_rtl(p)
            p.paragraph_format.space_before = Pt(15 if level == 1 else 11)
            p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            add_inline(p, text, size=13.5 if level == 2 else 17)
            i += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_rtl(p)
            p.paragraph_format.right_indent = Inches(0.28)
            p.paragraph_format.left_indent = Inches(0)
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, stripped[2:], size=11.2)
            i += 1
            continue
        paragraph_lines.append(stripped)
        i += 1
    flush_paragraph()


def keep_tables_together(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_together = True


def main():
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_markdown_body(doc, markdown)
    keep_tables_together(doc)
    setup_footer(doc)
    props = doc.core_properties
    props.title = "العرض الفني والتجاري لصيانة وتطوير تطبيق كنز الصحراء"
    props.subject = "فحص فني وخطة صيانة وتسعير"
    props.author = "م عمر"
    props.keywords = "كنز الصحراء صيانة Flutter WooCommerce دفع"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
