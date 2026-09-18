from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_client_problems_pdf import ISSUES, REMOVE_ROWS
from build_kanz_report import (
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_table_borders,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "word" / "تقرير_مشكلات_تطبيق_كنز_الصحراء_للعميل.docx"

BLACK = "171717"
MUTED = "5D5D5D"
GOLD = "8A6717"
RED = "B42318"
GREEN = "18794E"

ARABIC_ORDINALS = [
    "الأولى",
    "الثانية",
    "الثالثة",
    "الرابعة",
    "الخامسة",
    "السادسة",
    "السابعة",
    "الثامنة",
    "التاسعة",
    "العاشرة",
    "الحادية عشرة",
    "الثانية عشرة",
    "الثالثة عشرة",
    "الرابعة عشرة",
    "الخامسة عشرة",
    "السادسة عشرة",
    "السابعة عشرة",
    "الثامنة عشرة",
    "التاسعة عشرة",
    "العشرون",
    "الحادية والعشرون",
    "الثانية والعشرون",
    "الثالثة والعشرون",
    "الرابعة والعشرون",
    "الخامسة والعشرون",
    "السادسة والعشرون",
    "السابعة والعشرون",
    "الثامنة والعشرون",
    "التاسعة والعشرون",
    "الثلاثون",
    "الحادية والثلاثون",
    "الثانية والثلاثون",
    "الثالثة والثلاثون",
    "الرابعة والثلاثون",
    "الخامسة والثلاثون",
    "السادسة والثلاثون",
    "السابعة والثلاثون",
    "الثامنة والثلاثون",
]


def set_paragraph_rtl(paragraph, alignment=None):
    """Use Word's RTL paragraph mapping with a physical right edge."""
    # Microsoft Word mirrors left/right paragraph alignment when bidi is set.
    # LEFT therefore renders at the physical right margin for Arabic paragraphs.
    paragraph.alignment = alignment or WD_ALIGN_PARAGRAPH.LEFT
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.28


def set_run_rtl(run):
    """Mark Arabic runs as complex-script RTL so Word and PDF render identically."""
    r_pr = run._element.get_or_add_rPr()
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "1")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=9, color=MUTED)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after in [
        ("Title", 24, 0, 10),
        ("Heading 1", 18, 14, 8),
        ("Heading 2", 14.5, 11, 5),
    ]:
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLACK)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_rtl(footer, WD_ALIGN_PARAGRAPH.CENTER)
    run = footer.add_run("كنز الصحراء   تقرير المشكلات وخطة المعالجة   ")
    set_run_font(run, size=8.5, color=MUTED)
    set_run_rtl(run)
    add_page_number(footer)


def add_text(doc, text, *, bold=False, color=BLACK, size=11.5, after=7, keep=False):
    para = doc.add_paragraph()
    set_paragraph_rtl(para)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.25
    para.paragraph_format.keep_with_next = keep
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_run_rtl(run)
    return para


def add_heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    set_paragraph_rtl(para)
    run = para.add_run(text)
    set_run_font(run, size=18 if level == 1 else 14.5, bold=True, color=BLACK)
    set_run_rtl(run)
    return para


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("تقرير مشكلات تطبيق كنز الصحراء وخطة المعالجة")
    set_run_font(run, size=24, bold=True, color=BLACK)
    set_run_rtl(run)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    run = subtitle.add_run("شرح مبسط للإدارة وفريق العمل")
    set_run_font(run, size=13, color=MUTED)
    set_run_rtl(run)

    for _ in range(3):
        doc.add_paragraph()
    add_heading(doc, "الملخص التنفيذي", 1)
    add_text(doc, "راجعنا الملاحظات المسجلة في التقرير السابق، ثم قارنا التطبيق الحالي بموقع كنز الصحراء وصفحة الدفع الفعلية. اتضح أن التطبيق مبني على قالب متجر عام اشتراه المطور السابق، وما زالت داخله وظائف كثيرة لا يستخدمها المتجر.")
    add_text(doc, "المشكلة الأوضح الآن هي أن الموقع لا ينشر مجموعة الخدمات التي يعتمد عليها تطبيق الهاتف. لذلك قد ينجح تصفح المنتجات والأسعار، ثم يفشل الدفع أو الدخول أو تعديل الحساب أو حذف الحساب أو إرسال تقييم عند الخطوة الأخيرة. وكشف الفحص أيضاً أن مفاتيح المتجر الموجودة داخل التطبيق ما زالت فعالة وتصل إلى خدمات الطلبات والعملاء، ولذلك يجب إلغاؤها قبل نشر أي نسخة جديدة.")
    add_text(doc, "المعالجة ستبدأ بإعادة حلقة الربط الضرورية بين الموقع والتطبيق واختبار رحلة العميل كاملة، ثم تنظيف الأجزاء الزائدة التي جاءت مع القالب. النتيجة المطلوبة هي تطبيق يعكس موقع كنز الصحراء، ويلتزم بسياسة المتجر، ويعرض طرق الدفع المعتمدة، ويقدم تجربة عربية واضحة من دون خيارات لا تخص المتجر.")
    doc.add_page_break()


def add_issue(doc, number, title, problem, finding, outcome):
    add_heading(doc, f"المشكلة {ARABIC_ORDINALS[number - 1]}  {title}", 2)
    add_text(doc, "المشكلة", bold=True, color=RED, size=10.5, after=2, keep=True)
    add_text(doc, problem)
    add_text(doc, "حالات الحدوث والأسباب", bold=True, color=GOLD, size=10.5, after=2, keep=True)
    add_text(doc, finding)
    add_text(doc, "النتيجة بعد الإصلاح", bold=True, color=GREEN, size=10.5, after=2, keep=True)
    add_text(doc, outcome, after=10)


def add_removal_table(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_rtl(table)
    table.autofit = False
    widths = [Inches(1.55), Inches(2.65), Inches(2.45)]
    headers = ["المجموعة", "أمثلة موجودة في القالب", "الحالة المطلوبة"]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = widths[idx]
        set_cell_shading(cell, "8A6717")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, bottom=120, start=130, end=130)
        para = cell.paragraphs[0]
        set_paragraph_rtl(para)
        run = para.add_run(headers[idx])
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")
        set_run_rtl(run)
    set_repeat_table_header(table.rows[0])

    for row_no, row_data in enumerate(REMOVE_ROWS, 1):
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].width = widths[idx]
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx], top=110, bottom=110, start=130, end=130)
            if row_no % 2 == 0:
                set_cell_shading(cells[idx], "F7F1E3")
            para = cells[idx].paragraphs[0]
            set_paragraph_rtl(para)
            run = para.add_run(value)
            set_run_font(run, size=9.3, bold=(idx == 0), color=BLACK)
            set_run_rtl(run)
    set_table_borders(table)
    doc.add_paragraph()


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    add_cover(doc)
    add_heading(doc, "المشكلات والنتائج المتوقعة", 1)
    for index, data in enumerate(ISSUES, 1):
        if index == 37:
            doc.add_page_break()
        add_issue(doc, index, *data)

    doc.add_page_break()
    add_heading(doc, "الميزات الزائدة التي ستراجع للإزالة", 1)
    add_text(doc, "هذه الأجزاء جاءت مع القالب العام ولا تمثل وظائف متجر كنز الصحراء. ستراجع علاقتها بالدفع والشحن أولاً، ثم تزال من الواجهة وملف التطبيق عندما يثبت أنها غير مستخدمة.")
    add_removal_table(doc)

    add_heading(doc, "ما الذي سيلاحظه العميل بعد الإصلاح", 1)
    outcomes = [
        "إتمام الطلب يعمل ولا يبقى العميل أمام زر أو شاشة معلقة.",
        "المنتجات والأسعار والمخزون والأقسام متوافقة مع الموقع، وتظهر تغييرات العرض المعتمدة من دون انتظار تحديث كامل للتطبيق.",
        "طرق الدفع واضحة ومحددة وفق المتجر، ولا تظهر خيارات لا يستخدمها كنز الصحراء.",
        "التطبيق أخف في الوظائف والصلاحيات وأسهل في التحديث والصيانة، مع قياس فعلي للحجم والأداء قبل المقارنة وبعدها.",
        "رسائل الخطأ بالعربية توضح ما حدث وما الذي يمكن للعميل فعله، بدلاً من شاشة فارغة أو تحميل لا ينتهي.",
        "كل إصدار يمر باختبارات للشراء والدخول والتحديث والعربية، مع مراقبة الأعطال بعد النشر.",
    ]
    for idx, value in enumerate(outcomes, 1):
        add_text(doc, f"{value}", after=4)

    doc.core_properties.title = "تقرير مشكلات تطبيق كنز الصحراء وخطة المعالجة"
    doc.core_properties.author = "م عمر"
    doc.core_properties.subject = "شرح غير تقني للمشكلات والنتائج بعد الإصلاح"
    doc.save(OUTPUT)
    print("client report DOCX created")


if __name__ == "__main__":
    build()
